"""Перенаправить QR «Приложение» (слайд «Ссылки») на другой адрес и пересобрать PDF — без шаблона и полной сборки.

    .venv/bin/python presentation/set_app_qr.py http://192.168.1.5:8000

scripts/lan.sh вызывает его сам: IP ноутбука меняется вместе с сетью, а QR в презентации должен вести на него.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
from build_tbank import OUT, export_pdf  # noqa: E402

def app_qr(prs):
    """Слайд и картинка QR приложения: build_tbank называет её qr-app, номер слайда меняется вместе со сборкой."""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13 and shape.name == "qr-app":
                return slide, shape
    raise SystemExit("в презентации нет картинки qr-app — пересоберите её build_tbank.py")


def main(url: str) -> None:
    png = Path(tempfile.mkdtemp()) / "qr_app.png"
    subprocess.run(["node", "-e", "require('qrcode').toFile(process.argv[1], process.argv[2], "
                    "{width: 800, margin: 1, color: {dark: '#333333', light: '#FFFFFF'}})", str(png), url],
                   cwd=Path(__file__).parent, check=True)
    prs = Presentation(str(OUT))
    slide, pic = app_qr(prs)
    rid = pic._element.blipFill.blip.rEmbed
    slide.part.related_part(rid)._blob = png.read_bytes()
    prs.save(str(OUT))
    print("QR приложения →", url)
    export_pdf(OUT)


if __name__ == "__main__":
    if len(sys.argv) != 2 or not sys.argv[1].startswith(("http://", "https://")):
        sys.exit("укажите адрес: set_app_qr.py http://<IP>:8000")
    main(sys.argv[1])
