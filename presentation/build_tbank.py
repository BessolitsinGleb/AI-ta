"""Презентация в фирменном шаблоне Т-Банк Образование.

python presentation/build_tbank.py ШАБЛОН.pptx [АССЕТЫ]
→ presentation/Amba_TBank.pptx

Структура — по методичке «Как презентовать проект на защите» (7 минут): цель (2 мин) → ход работы по этапам
с отсылкой к плану (4 мин) → результат с оценкой качества и честными ограничениями (1 мин).
Берём слайды-примеры из шаблона (один и тот же макет можно взять несколько раз), заполняем, остальное удаляем.
Тексты выступления берутся из docs/SPEECH.md и кладутся в заметки к слайдам.
"""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(sys.argv[1])
ASSETS = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "presentation" / "assets"
OUT = ROOT / "presentation" / "Amba_TBank.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

APP_URL = "https://amba-qm59.onrender.com"

# Порядок: номер слайда-примера в шаблоне → наш слайд (68 — этапы работы, берётся трижды).
ORDER = [13, 24, 64, 20, 15, 68, 68, 68, 59, 21, 53, 26, 67, 35, 82]
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


# ---------- текст с сохранением фирменного форматирования ----------

def _clean_runs(p: etree._Element) -> etree._Element:
    runs = p.findall(f"{A}r")
    for extra in runs[1:]:
        p.remove(extra)
    for tag in ("br", "fld"):
        for el in p.findall(f"{A}{tag}"):
            p.remove(el)
    if not runs:
        r = etree.SubElement(p, f"{A}r")
        end = p.find(f"{A}endParaRPr")
        rpr = copy.deepcopy(end) if end is not None else etree.Element(f"{A}rPr")
        rpr.tag = f"{A}rPr"
        r.append(rpr)
        etree.SubElement(r, f"{A}t")
        if end is not None:
            p.remove(end)
            p.append(end)
    return p


def set_text(shape_or_tf, paras: list[str] | str) -> None:
    """Заменить текст, сохранив оформление: i-й абзац копирует стиль i-го абзаца шаблона (или последнего)."""
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, "text_frame") else shape_or_tf
    paras = [paras] if isinstance(paras, str) else paras
    body = tf._txBody
    old = body.findall(f"{A}p")
    protos = [copy.deepcopy(p) for p in old] or [etree.Element(f"{A}p")]
    for p in old:
        body.remove(p)
    for i, text in enumerate(paras):
        p = _clean_runs(copy.deepcopy(protos[min(i, len(protos) - 1)]))
        r = p.find(f"{A}r")
        r.find(f"{A}t").text = text
        end = p.find(f"{A}endParaRPr")
        if end is not None:
            p.remove(end)
            p.append(end)
        body.append(p)


def shape(slide, shape_id: int):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
        if sh.shape_type == 6:
            for c in sh.shapes:
                if c.shape_id == shape_id:
                    return c
    raise KeyError(shape_id)


def put(slide, shape_id: int, paras) -> None:
    set_text(shape(slide, shape_id), paras)


def remove(slide, shape_id: int) -> None:
    el = shape(slide, shape_id)._element
    el.getparent().remove(el)


def picture(slide, path: Path, x: float, y: float, w: float | None = None, h: float | None = None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def duplicate(prs, src):
    """Копия слайда-примера в конец презентации: фигуры, картинки (связи пересчитываются), фон."""
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(("/notesSlide", "/slideLayout")):
            continue
        rid_map[rid] = (dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else dst.part.relate_to(rel._target, rel.reltype))
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}nvGrpSpPr", "}grpSpPr")):
            continue
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr, val in list(node.attrib.items()):
                if attr.startswith(R) and val in rid_map:
                    node.set(attr, rid_map[val])
        dst.shapes._spTree.append(new)
    bg = src._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(bg))
    return dst


def stage_progress(slide, k: int) -> None:
    """Полоса «01 02 03 04» внизу слайда этапа: заполнена до текущего этапа, его кружок — фирменный жёлтый."""
    dots = [shape(slide, i) for i in (8, 10, 11, 12)]
    on = copy.deepcopy(dots[0]._element.spPr.find(f"{A}solidFill"))
    off = copy.deepcopy(dots[1]._element.spPr.find(f"{A}solidFill"))
    for i, dot in enumerate(dots, 1):
        sp = dot._element.spPr
        sp.replace(sp.find(f"{A}solidFill"), copy.deepcopy(on if i <= k else off))
    bar = shape(slide, 7)
    bar.width = dots[k - 1].left + dots[k - 1].width // 2 - bar.left


def stage(slide, k: int, title: str, blocks: list[tuple[str, str]]) -> None:
    """Слайд этапа по методичке: что и зачем → инструменты → трудность → результат."""
    put(slide, 5, f"Этап {k} из 4 · {title}")
    for (head, text), (hid, tid) in zip(blocks, [(29, 28), (32, 31), (35, 34), (38, 37)]):
        put(slide, hid, head)
        put(slide, tid, text)
        body = shape(slide, tid)
        body.height = Inches(3.7)  # до низа карточки: текст крупнее и не теснится
        for box, size in ((shape(slide, hid), 24), (body, 19)):
            for para in box.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(size)
    stage_progress(slide, k)


def replace_photo(slide, shape_id: int, path: Path) -> None:
    """Фото-пример шаблона → наша картинка в той же рамке: обрезка «по заполнению», форма рамки (скругления) та же."""
    from PIL import Image
    old = shape(slide, shape_id)
    x, y, w, h = old.left, old.top, old.width, old.height
    geom = old._element.spPr.find(f"{A}prstGeom")
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    iw, ih = Image.open(path).size
    frame, img = w / h, iw / ih
    if img > frame:
        pic.crop_left = pic.crop_right = (1 - frame / img) / 2
    else:
        pic.crop_top = pic.crop_bottom = (1 - img / frame) / 2
    if geom is not None:
        pic._element.spPr.replace(pic._element.spPr.find(f"{A}prstGeom"), copy.deepcopy(geom))
    old._element.addprevious(pic._element)  # на место примера, под остальными фигурами
    old._element.getparent().remove(old._element)


def fill_chart(frame, cats: list[str], values: list[float], fmt: str) -> None:
    from pptx.chart.data import CategoryChartData
    data = CategoryChartData()
    data.categories = cats
    data.add_series("Ряд 1", values)
    frame.chart.replace_data(data)
    labels = frame.chart.plots[0].data_labels
    labels.number_format, labels.number_format_is_linked = fmt, False
    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    for ser in frame.chart._chartSpace.iter(f"{C}ser"):
        for d in ser.findall(f"{C}dPt"):  # раскраска точек из примера (минус — красным) нам не нужна
            ser.remove(d)
        inv = ser.find(f"{C}invertIfNegative")
        if inv is not None:
            inv.set("val", "0")
        for nf in ser.iter(f"{C}numFmt"):
            nf.set("formatCode", fmt)
            nf.set("sourceLinked", "0")
        dl = ser.find(f"{C}dLbls")
        if dl is not None and dl.find(f"{C}numFmt") is None:  # формат подписей задаём на уровне ряда
            dl.insert(0, etree.Element(f"{C}numFmt", formatCode=fmt, sourceLinked="0"))
    from pptx.dml.color import RGBColor
    series = frame.chart.plots[0].series[0]
    for i in range(len(values)):
        pt = series.points[i].format.fill
        pt.solid()
        pt.fore_color.rgb = RGBColor(0xFF, 0xDD, 0x2D) if i == len(values) - 1 else RGBColor(0xE0, 0xE1, 0xE3)


# ---------- таблицы ----------

def shrink_table(table, rows: int, cols: int) -> None:
    tbl = table._tbl
    grid = tbl.tblGrid
    total = sum(int(gc.get("w")) for gc in grid.findall(f"{A}gridCol"))
    for gc in grid.findall(f"{A}gridCol")[cols:]:
        grid.remove(gc)
    for tr in tbl.findall(f"{A}tr")[rows:]:
        tbl.remove(tr)
    for tr in tbl.findall(f"{A}tr"):
        for tc in tr.findall(f"{A}tc")[cols:]:
            tr.remove(tc)
    return total


def set_col_widths(table, inches: list[float]) -> None:
    for gc, w in zip(table._tbl.tblGrid.findall(f"{A}gridCol"), inches):
        gc.set("w", str(int(Inches(w))))


# ---------- данные ----------

def speech() -> list[tuple[str, str]]:
    md = (ROOT / "docs" / "SPEECH.md").read_text(encoding="utf-8")
    return re.findall(r"^## \d+\. (.+? — \d:\d\d)\n\n(.+?)\n", md, flags=re.M)


def evals() -> list[dict]:
    return [json.loads((ROOT / "eval" / "results" / f"{v}_core10.json").read_text())["aggregate"] for v in ("v1", "v2", "v3")]


def avatar(path: Path, initials: str, dark: bool) -> None:
    bg, fg = ("#333333", "#FFDD2D") if dark else ("#FFDD2D", "#333333")
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="400" height="400"><rect width="400" height="400" rx="80" fill="{bg}"/>'
           f'<text x="200" y="262" font-family="Arial, Helvetica, sans-serif" font-weight="700" font-size="170" '
           f'text-anchor="middle" fill="{fg}">{initials}</text></svg>')
    src = path.with_suffix(".svg")
    src.write_text(svg)
    subprocess.run(["node", "-e", f"require('sharp')('{src}').png().toFile('{path}')"], cwd=ROOT / "presentation", check=True)


# ---------- сборка ----------

def main() -> None:
    prs = Presentation(str(TEMPLATE))
    sld_ids = list(prs.slides._sldIdLst)
    src = list(prs.slides)
    keep, used = [], set()
    for n in ORDER:
        if n in used:  # тот же макет ещё раз — копия слайда-примера
            duplicate(prs, src[n - 1])
            keep.append(prs.slides._sldIdLst[-1])
        else:
            keep.append(sld_ids[n - 1])
            used.add(n)
    for el in list(prs.slides._sldIdLst):
        if el not in keep:
            prs.part.drop_rel(el.rId)
        prs.slides._sldIdLst.remove(el)
    for el in keep:
        prs.slides._sldIdLst.append(el)
    S = list(prs.slides)
    v1, v2, v3 = evals()

    # 1. Титул
    s = S[0]
    put(s, 4, "Амба — AI-гид по Приморью")
    put(s, 5, "Агент собирает план поездки и проверяет, что он сходится с реальностью")
    remove(s, 16)
    picture(s, ASSETS / "amba.png", 12.6, 1.4, h=9.6)

    # 2. Проблема и для кого
    s = S[1]
    put(s, 6, "Какую проблему решаем")
    put(s, 14, "Данные разбросаны")
    put(s, 4, ["Погода, расписания, режим музеев, АЗС, жильё, кафе — всё в разных местах.",
               "Чат-боты строят план, но не проверяют его: ведут в закрытый музей или на велосипед под ливень.",
               "В Приморье цена ошибки выше: сопки, грунтовки, редкие рейсы и паромы."])
    put(s, 10, "Для кого")
    put(s, 12, "Кому полезна Амба")
    put(s, 11, ["Гости края, которые впервые во Владивостоке",
                "Местные жители: «куда на выходные»",
                "Travel-сервисы банков и агрегаторы: план → бронь → оплата",
                "Туристические центры и турбазы края"])

    # ===== Часть 1. Цель проекта (2 минуты): мотивация → исследование аналогов → цель =====

    # 3. Исследование: аналоги по критериям, на которых ломается поездка
    s = S[2]
    put(s, 3, "Как это решают сейчас")
    tbl = shape(s, 5).table
    total = shrink_table(tbl, rows=7, cols=5)
    rows = [["Что нужно для поездки", "ChatGPT, Алиса", "Туту, Островок", "Путеводители", "Амба"],
            ["Запрос обычными словами", "да", "нет", "нет", "да"],
            ["Весь план: места, путь, ночлег, еда", "да", "частично", "частично", "да"],
            ["Погода, режим работы, расписания", "нет", "частично", "частично", "да"],
            ["Транспорт: АЗС, расстояния, рейсы", "нет", "частично", "нет", "да"],
            ["Работает ли кафе или музей сейчас", "нет", "нет", "нет", "да"],
            ["Ссылки, чтобы сразу оформить", "нет", "да", "частично", "да"]]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_text(tbl.cell(r, c).text_frame, val)
    w = total / 914400
    set_col_widths(tbl, [w * 0.37, w * 0.15, w * 0.16, w * 0.17, w * 0.15])
    from pptx.dml.color import RGBColor
    # Палитра «положительные / отрицательные значения» из шаблона: да — зелёный, частично — жёлтый, нет — розовый.
    tone = {"да": RGBColor(0xDA, 0xEC, 0xBD), "частично": RGBColor(0xFF, 0xEE, 0x96), "нет": RGBColor(0xFF, 0xE1, 0xD5)}
    for r in range(7):
        tbl.rows[r].height = Inches(7.9 / 7)
        for c in range(5):
            cell = tbl.cell(r, c)
            if r > 0 and c > 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = tone[rows[r][c]]
            if r == 0 and c == 4:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xDD, 0x2D)
            for para in cell.text_frame.paragraphs:
                if c > 0:
                    para.alignment = 2  # по центру: цветные ячейки читаются как значки
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    run.font.bold = (c == 4) or r == 0

    # 4. Цель — вытекает из сравнения: никто не проверяет план по реальным данным
    s = S[3]
    put(s, 5, "Цель: AI-гид, который по запросу обычными словами собирает план поездки по Приморью "
              "и сам проверяет его по реальным данным")
    put(s, 6, "Цель проекта · следует из сравнения: аналоги не проверяют, сходится ли план с реальностью")
    goal, note = shape(s, 5), shape(s, 6)
    goal.top, goal.height = Inches(3.0), Inches(3.4)
    note.top = Inches(6.9)
    for para in goal.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(44)

    # ===== Часть 2. Ход работы (4 минуты): задачи и слайд на каждый этап =====

    # 5. Задачи — из цели; дальше каждый этап решает свою задачу («Этап N из 4» и полоса внизу)
    s = S[4]
    put(s, 3, ["Задачи"])
    for sid, text in zip((12, 13, 14, 15), ("1. Подключить данные", "2. Построить агентный цикл",
                                            "3. Научить проверять план", "4. Замерить три версии")):
        put(s, sid, text)
    replace_photo(s, 11, ROOT / "docs" / "img" / "map.png")

    # 6–8. Этапы: что и зачем → инструменты → трудность → результат (в числах)
    stage(S[5], 1, "Данные и инструменты", [
        ("Что и зачем", "Агенту нужны факты, а не догадки: погода, дороги, кафе, заправки, расписания"),
        ("Инструменты", "Open-Meteo, OSRM, OpenStreetMap, Википедия — открытые и бесплатные"),
        ("Трудность", "Overpass отвечал 40 секунд, а OSM помнит закрытые кафе. Решение: снимок всех АЗС и проверка сайтов"),
        ("Результат", "10 тулов и 5 источников. Сервис упал — тул берёт запасной, а не падает")])
    stage(S[6], 2, "Агентный цикл", [
        ("Что и зачем", "Один промпт выдумывает. Агент сам решает, какие данные ему нужны, и зовёт тулы"),
        ("Инструменты", "Свой цикл на FastAPI, tool calling, DeepSeek V4.1 Flash через LLM Proxy"),
        ("Трудность", "Модель ходила по кругу и обрезала длинный план. Решение: лимит шагов и сжатие контекста"),
        ("Результат", "1–2 уточнения, тулы параллельно, план за минуту, правки прямо в чате")])
    stage(S[7], 3, "Проверка плана", [
        ("Что и зачем", "План должен сходиться с реальностью. Проверяет код, а не модель"),
        ("Инструменты", "15 проверок: сроки, транспорт, километры, погода, режим работы, АЗС, рейсы"),
        ("Трудность", "Обед за 40 км и закрытый бар. Решение: «еда рядом» и проверка сайтов заведений"),
        ("Результат", "Противоречие — агент исправляет сам. Не про Приморье — отказ без модели")])

    # 9. Этап 4: замеры трёх версий — графиком, а не списком чисел
    s = S[8]
    put(s, 3, "Этап 4 из 4 · Замеры трёх версий")
    put(s, 11, "10 запросов: 7 поездок разным транспортом и 3 нецелевых. Качество — верные тулы и доля "
               "пройденных проверок. Токены и цена — из ответа провайдера модели")
    put(s, 5, "Качество, от 0 до 1")
    put(s, 6, "Токены на запрос, тыс.")
    versions_ru = ["v1 базовая", "v2 сжатие и сверка", "v3 финальная"]
    fill_chart(shape(s, 7), versions_ru, [v1["quality"], v2["quality"], v3["quality"]], "0.00")
    fill_chart(shape(s, 8), versions_ru, [round(v["tokens_avg"] / 1000) for v in (v1, v2, v3)], "0")
    put(s, 10, f"{round(v3['converge_rate'] * 100)}%")
    put(s, 9, f"планов сходятся, было {round(v1['converge_rate'] * 100)}%")
    big, cap = shape(s, 10), shape(s, 9)
    big.left, big.width = Inches(16.9), Inches(3.4)
    cap.left, cap.width, cap.top = Inches(16.9), Inches(3.4), Inches(10.35)

    # 10. Что нашли замеры — трудности этапа 4
    s = S[9]
    put(s, 6, "Замеры нашли настоящие баги агента — все исправлены в финальной версии")
    put(s, 5, "3")
    put(s, 7, ["Длинный план обрезался по лимиту токенов",
               "Модель присылала дни строкой, а не списком",
               "На невыполнимом запросе агент ходил по кругу"])
    for para in shape(s, 7).text_frame.paragraphs:
        para.space_after = Pt(18)

    # ===== Часть 3. Результат (1 минута): что получилось и насколько решает задачу =====

    # 11. Что получилось
    s = S[10]
    put(s, 4, "Результат")
    put(s, 2, "Что получилось")
    put(s, 7, "Запрос обычными словами → план по дням, карта и адреса")
    put(s, 6, "У каждого места отметка: работает ли оно и как это проверено")
    put(s, 8, "Правки в чате, аккаунты, история поездок, идеи по анкете")
    shape(s, 13).insert_picture(str(ROOT / "docs" / "img" / "plan.png"))

    # 12. Решает ли задачу — и честные ограничения
    s = S[11]
    put(s, 3, "Решает ли задачу")
    put(s, 2, "Да: план за минуту вместо десятка сайтов, и он сходится с реальностью. Но не всё можно "
              "проверить — ограничения показываем честно")
    put(s, 10, "Что получилось")
    put(s, 8, [f"{round(v3['converge_rate'] * 100)}% планов в замерах проходят все проверки",
               f"Токенов в {v1['tokens_avg'] / v3['tokens_avg']:.1f} раза меньше, чем у базовой версии".replace(".", ","),
               "Закрытые кафе отсеиваются по их же сайтам",
               "Аналоги план не проверяют — Амба проверяет"])
    put(s, 11, "Ограничения")
    put(s, 9, ["Остатки топлива и свободные номера никто не публикует",
               "У многих кафе нет сайта — пишем «уточните»",
               "Пригородные расписания — наш набор данных",
               "Бесплатный хостинг засыпает: первый ответ ~1 мин"])

    # ===== Команда, ссылки, финал =====

    # 13. Команда
    s = S[12]
    put(s, 5, "Наша команда")
    put(s, 6, "Семь участников: люди и AI-агенты, у каждого своя зона ответственности")
    for sid, w in ((5, 14.0), (6, 17.0)):  # шире, чтобы в любом шрифте оставались в одну строку
        box = shape(s, sid)
        box.left, box.width = Inches((21.65 - w) / 2), Inches(w)
    team = [  # (группа, плейсхолдер фото, имя, роль, инициалы, AI)
        (54, 3, "Opus 5.5", "Team Lead", "O5", True),
        (11, 10, "Саша Дё", "Backend developer", "СД", False),
        (13, 4, "Глеб Бессолицин", "Backend developer", "ГБ", False),
        (18, 14, "Анфиса Кирсанова", "ML / Frontend", "АК", False),
        (39, 12, "Ксюша Яковлева", "ML / Frontend", "КЯ", False),
        (44, 7, "Арнольд Гончаров", "QA-engineer", "АГ", False),
        (49, 17, "Luna 6", "Debug-engineer", "L6", True),
    ]
    for gid, ph, name, role, ini, ai in team:
        g = shape(s, gid)
        name_box, role_box = list(g.shapes)
        set_text(name_box, name.split(" ", 1) if " " in name and not ai else [name])
        set_text(role_box, role)
        for box, size in ((name_box, 26), (role_box, 16)):  # с запасом под любой шрифт: длинные фамилии и роли
            for para in box.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(size)
        img = ASSETS / f"avatar_{ini}.png"
        avatar(img, ini, ai)
        shape(s, ph).insert_picture(str(img))

    # 14. Сайт команды + приложение: две одинаковые QR-плашки
    s = S[13]
    put(s, 9, ["Сайт", "команды"])
    for para in shape(s, 9).text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(56)  # было 72: с запасом под широкие шрифты
    remove(s, 10)  # без ссылок на слайде: только QR-коды
    put(s, 5, "Приложение")
    put(s, 2, "Сканируйте QR-код")
    put(s, 4, "и откройте Амбу")
    shape(s, 2).width = shape(s, 4).width = shape(s, 5).width = Inches(4.6)
    put(s, 15, "14")  # в этом примере номер слайда — обычный текст
    plate = shape(s, 16)
    set_text(plate, " ")
    twin = copy.deepcopy(plate._element)  # такая же фирменная плашка под QR, но в белой карточке
    twin_id = max(int(e.get("id")) for e in s.shapes._spTree.iter() if e.tag.endswith("}cNvPr")) + 1
    twin.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", str(twin_id))
    twin.find(f".//{A}off").set("x", str(int(Inches(8.3))))
    plate._element.addnext(twin)
    twin_shape = shape(s, twin_id)
    twin_shape.fill.solid()
    from pptx.dml.color import RGBColor
    twin_shape.fill.fore_color.rgb = RGBColor(0xEF, 0xF0, 0xF1)
    for qr_png, left, name in ((ASSETS / "qr_site.png", 8.3, "qr-site"), (ASSETS / "qr_demo.png", plate.left / 914400, "qr-app")):
        picture(s, qr_png, left + 0.45, plate.top / 914400 + 0.45, w=plate.width / 914400 - 0.9).name = name  # qr-app меняет set_app_qr.py

    # 15. Финал
    s = S[14]
    put(s, 6, ["Спасибо", "за внимание!"])
    thanks = shape(s, 6)
    thanks.top, thanks.height = Inches(6.6), Inches(3.6)
    for para in thanks.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(92)  # было 115: в широких шрифтах «внимание!» обрезалось
    picture(s, ASSETS / "qr_egg.png", 9.9, 1.3, w=3.3)
    picture(s, ASSETS / "amba.png", 10.25, 4.9, h=2.2)

    # Заметки с текстом выступления и таймингом.
    for slide, (title, text) in zip(S, speech()):
        slide.notes_slide.notes_text_frame.text = f"[{title}] {text}"

    prs.save(str(OUT))
    print("saved", OUT, "slides:", len(prs.slides))
    export_pdf(OUT)


PDF_FONTS = {b"T-Sans Medium": b"PT Sans Caption", b"TINKOFFSANS-MEDIUM": b"PT Sans Caption", b"T-Sans": b"PT Sans", b"Tinkoff Sans": b"PT Sans"}


def export_pdf(pptx: Path) -> Path:
    """PDF с зашитым шрифтом: фирменный T-Sans не у всех установлен, в PDF подставляем PT Sans — он есть в macOS и открыт (OFL)."""
    import shutil
    import tempfile
    import zipfile

    work = Path(tempfile.mkdtemp())
    tmp = work / pptx.name
    with zipfile.ZipFile(pptx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                for old, new in PDF_FONTS.items():
                    data = data.replace(b'typeface="' + old + b'"', b'typeface="' + new + b'"')
            zout.writestr(item, data)
    profile = work / "lo"
    (profile / "user" / "fonts").mkdir(parents=True)
    pt_sans = Path("/System/Library/Fonts/Supplemental/PTSans.ttc")
    if pt_sans.exists():
        shutil.copy(pt_sans, profile / "user" / "fonts")  # LibreOffice видит только свои шрифты и шрифты профиля
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(tmp)],
                   check=True, capture_output=True, timeout=600)
    out = pptx.with_suffix(".pdf")
    shutil.copy(work / pptx.with_suffix(".pdf").name, out)
    print("saved", out)
    return out


if __name__ == "__main__":
    main()
